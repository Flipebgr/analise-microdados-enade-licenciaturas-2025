from __future__ import annotations

import csv
import io
import re
from pathlib import Path
from zipfile import ZipFile

import openpyxl
import pytest

from src.edicoes import ENADE_2017, ENADE_2025_LICENCIATURAS
from src.extracao.inventario import inventariar_arquivos
from src.utilitarios.leitura import obter_cursos_area

pytestmark = pytest.mark.integration

ROOT = Path(__file__).resolve().parents[2]


@pytest.mark.parametrize(
    ("contrato", "caminho"),
    [
        (ENADE_2017, ROOT / "dados_brutos" / "enade_2017" / "microdados_enade_2017_LGPD.zip"),
        (
            ENADE_2025_LICENCIATURAS,
            ROOT / "dados_brutos" / "microdados_enade_licenciaturas_2025.zip",
        ),
    ],
)
def test_inventario_e_headers_correspondem_ao_contrato(contrato, caminho):
    if not caminho.exists():
        pytest.skip(f"Fonte local ausente: {caminho}")

    padrao = re.compile(rf"{re.escape(contrato.prefixo_arquivos)}(\d+)\.txt$")
    with ZipFile(caminho) as arquivo_zip:
        encontrados = {}
        for nome in arquivo_zip.namelist():
            correspondencia = padrao.search(nome)
            if correspondencia:
                encontrados[int(correspondencia.group(1))] = nome

        assert sorted(encontrados) == list(range(1, contrato.quantidade_arquivos + 1))

        for numero, nome in encontrados.items():
            with arquivo_zip.open(nome) as arquivo:
                leitor = csv.reader(io.TextIOWrapper(arquivo, encoding="utf-8-sig"), delimiter=";")
                cabecalho = set(next(leitor))
            obrigatorias = set(contrato.schema_arquivo(numero).colunas_obrigatorias)
            assert obrigatorias <= cabecalho, f"Colunas ausentes em {Path(nome).name}"


@pytest.mark.parametrize(
    ("contrato", "caminho"),
    [
        (
            ENADE_2017,
            ROOT / "dados_brutos" / "enade_2017" / "resultados_conceito_enade_2017.xlsx",
        ),
        (
            ENADE_2025_LICENCIATURAS,
            ROOT / "dados_brutos" / "conceito_enade_licenciaturas.xlsx",
        ),
    ],
)
def test_planilha_conceito_corresponde_ao_contrato(contrato, caminho):
    if not caminho.exists():
        pytest.skip(f"Fonte local ausente: {caminho}")

    workbook = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    try:
        assert contrato.conceito.aba in workbook.sheetnames
        worksheet = workbook[contrato.conceito.aba]
        cabecalho = {
            valor
            for valor in next(worksheet.iter_rows(values_only=True))
            if isinstance(valor, str) and valor
        }
        colunas_origem = {origem for origem, _ in contrato.conceito.mapa_colunas}
        assert colunas_origem <= cabecalho
    finally:
        workbook.close()


def test_piloto_2017_localiza_curso_ufpa_sem_extrair_zip():
    caminho = ROOT / "dados_brutos" / "enade_2017" / "microdados_enade_2017_LGPD.zip"
    if not caminho.exists():
        pytest.skip(f"Fonte local ausente: {caminho}")

    inventario = inventariar_arquivos(caminho, ENADE_2017)
    cursos = obter_cursos_area(caminho, ENADE_2017, 1601, chunksize=10_000)

    assert len(inventario) == 42
    assert "12027" in cursos
