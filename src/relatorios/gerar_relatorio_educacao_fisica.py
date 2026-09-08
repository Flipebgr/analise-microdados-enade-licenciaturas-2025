from __future__ import annotations

# ruff: noqa: E501

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from src.educacao_fisica import EDUCACAO_FISICA
from src.educacao_fisica.rotulos_questionario import ROTULOS_DIMENSOES
from src.relatorios.conversao_pdf import converter_docx_para_pdf
from src.relatorios.figuras_relatorio import adicionar_figura
from src.relatorios.formatacao_abnt import configurar_cabecalho_rodape, configurar_documento
from src.relatorios.referencias import adicionar_referencias
from src.relatorios.resultado_relatorio import ResultadoRelatorio
from src.relatorios.tabelas_relatorio import adicionar_tabela

FONTE_DADOS = (
    "Elaboração própria com base nos microdados do Enade das Licenciaturas 2025 "
    "e na planilha de Conceito Enade."
)


def _ler_csv(path: Path) -> pd.DataFrame:
    if not path.exists():
        raise FileNotFoundError(
            f"Produto analítico ausente: {path}. Execute base e validação antes do relatório."
        )
    return pd.read_csv(path, low_memory=False)


def carregar_produtos(root: Path) -> dict[str, pd.DataFrame]:
    pasta = root / "dados_processados" / "educacao_fisica"
    return {
        "base": _ler_csv(pasta / "base_analitica_validada.csv"),
        "benchmark": _ler_csv(pasta / "benchmark_sensibilidade_resumo.csv"),
        "associacoes": _ler_csv(pasta / "associacoes_ecologicas.csv"),
        "consistencia": _ler_csv(pasta / "consistencia_dimensoes_processo.csv"),
        "itens": _ler_csv(pasta / "catalogo_itens_processo_formativo.csv"),
    }


def _num(valor) -> float:
    return pd.to_numeric(pd.Series([valor]), errors="coerce").iloc[0]


def _fmt(valor, casas: int = 2) -> str:
    valor = _num(valor)
    if pd.isna(valor):
        return "—"
    return f"{valor:.{casas}f}".replace(".", ",")


def _pct(valor, casas: int = 1) -> str:
    valor = _num(valor)
    if pd.isna(valor):
        return "—"
    return f"{100 * valor:.{casas}f}%".replace(".", ",")


def _int(valor) -> str:
    valor = _num(valor)
    if pd.isna(valor):
        return "—"
    return str(int(round(valor)))


def _ufpa(base: pd.DataFrame) -> pd.DataFrame:
    return base.loc[base["CO_IES"].eq(EDUCACAO_FISICA.co_ies_focal)].sort_values(
        "ROTULO_OFERTA"
    )


def _oferta(base: pd.DataFrame, co_curso: int) -> pd.Series:
    alvo = base.loc[pd.to_numeric(base["CO_CURSO"], errors="coerce").eq(co_curso)]
    if len(alvo) != 1:
        raise ValueError(f"Esperada uma oferta CO_CURSO={co_curso}; encontradas {len(alvo)}")
    return alvo.iloc[0]


def tabela_ofertas(base: pd.DataFrame) -> pd.DataFrame:
    u = _ufpa(base).copy()
    return pd.DataFrame(
        {
            "CO_CURSO": u["CO_CURSO"].astype(int),
            "Oferta": u["ROTULO_OFERTA"],
            "Inscritos": u["INSCRITOS_NUM"].round().astype("Int64"),
            "Participantes": u["PARTICIPANTES_NUM"].round().astype("Int64"),
            "Participação": u["TAXA_PARTICIPACAO_OFICIAL"].map(_pct),
            "Proficiência": u["PCT_PADRAO_PROFICIENCIA_NUM"].map(_pct),
            "Conceito": u["CONCEITO_ENADE_NUM"].round().astype("Int64"),
            "NT_GER": u["nt_ger_mean"],
            "NT_OBJ": u["nt_obj_mean"],
            "NT_DIS": u["nt_dis_mean"],
        }
    )


def tabela_perfil(base: pd.DataFrame) -> pd.DataFrame:
    bel = _oferta(base, 104598)
    cas = _oferta(base, 21849)
    linhas = [
        ("Sexo feminino", "sexo_feminino_pct", "sexo_n_valido", False),
        ("Idade média (anos)", "idade_media", "idade_n", True),
        ("Mãe com superior", "mae_superior_pct", "mae_superiorn_valido", False),
        ("Pai com superior", "pai_superior_pct", "pai_superiorn_valido", False),
        ("Renda até 3 SM", "renda_ate_3sm_pct", "renda_ate_3smn_valido", False),
        ("Trabalha", "trabalha_pct", "trabalhan_valido", False),
        ("Ação afirmativa", "acao_afirmativa_pct", "acao_afirmativan_valido", False),
        ("Auxílio permanência", "auxilio_permanencia_pct", "auxilio_permanencian_valido", False),
        ("Bolsa acadêmica", "bolsa_academica_pct", "bolsa_academican_valido", False),
        ("Estudo ≥4h/semana", "estudo_4h_ou_mais_pct", "estudo_4h_ou_maisn_valido", False),
        ("Pretende magistério", "pretende_magisterio_pct", "pretende_magisterion_valido", False),
    ]
    dados = []
    for rotulo, col, ncol, numero in linhas:
        dados.append(
            {
                "Indicador": rotulo,
                "Belém": _fmt(bel[col]) if numero else _pct(bel[col]),
                "N Belém": _int(bel[ncol]),
                "Castanhal": _fmt(cas[col]) if numero else _pct(cas[col]),
                "N Castanhal": _int(cas[ncol]),
            }
        )
    return pd.DataFrame(dados)


def tabela_dimensoes(base: pd.DataFrame) -> pd.DataFrame:
    bel = _oferta(base, 104598)
    cas = _oferta(base, 21849)
    dados = []
    for chave, rotulo in ROTULOS_DIMENSOES.items():
        col = f"dim_{chave}_media"
        dados.append(
            {
                "Dimensão": rotulo,
                "Belém": bel[col],
                "Castanhal": cas[col],
                "Dif. Castanhal-Belém": cas[col] - bel[col],
            }
        )
    return pd.DataFrame(dados)


def tabela_benchmark(benchmark: pd.DataFrame) -> pd.DataFrame:
    b = benchmark.loc[
        benchmark["CRITERIO"].eq("porte_50pct")
        & benchmark["INDICADOR"].eq("nt_ger_mean")
    ].copy()
    return b[
        [
            "ROTULO_ALVO",
            "N_COMPARAVEIS",
            "VALOR_ALVO",
            "MEDIA_BENCHMARK",
            "MEDIANA_BENCHMARK",
            "DIF_MEDIA",
            "DIF_MEDIANA",
        ]
    ].rename(
        columns={
            "ROTULO_ALVO": "Oferta",
            "N_COMPARAVEIS": "N comparáveis",
            "VALOR_ALVO": "NT_GER alvo",
            "MEDIA_BENCHMARK": "Média benchmark",
            "MEDIANA_BENCHMARK": "Mediana benchmark",
            "DIF_MEDIA": "Dif. média",
            "DIF_MEDIANA": "Dif. mediana",
        }
    )


def _adicionar_texto(doc: Document, texto: str) -> None:
    doc.add_paragraph(texto)


def _capa(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(72)
    p.add_run("UNIVERSIDADE FEDERAL DO PARÁ").bold = True
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.add_run("ENADE DAS LICENCIATURAS 2025").bold = True
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(72)
    r = p.add_run("EDUCAÇÃO FÍSICA")
    r.bold = True
    r.font.size = Pt(16)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.add_run(
        "Desempenho, perfil discente, processo formativo e benchmarks das ofertas da UFPA"
    ).bold = True
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(150)
    p.add_run("Belém\n2026")
    doc.add_page_break()


def gerar_docx(root: Path, produtos: dict[str, pd.DataFrame], destino: Path) -> None:
    base = produtos["base"]
    benchmark = produtos["benchmark"]
    associacoes = produtos["associacoes"]
    bel = _oferta(base, 104598)
    cas = _oferta(base, 21849)
    figuras = root / "figuras" / "educacao_fisica"

    doc = Document()
    configurar_documento(doc)
    configurar_cabecalho_rodape(doc)
    _capa(doc)

    doc.add_heading("RESUMO", level=1)
    p = doc.add_paragraph(style="Resumo")
    p.add_run(
        "Este relatório técnico-científico analisa 406 cursos de Educação Física no Enade "
        "das Licenciaturas 2025, com foco nas duas ofertas presenciais da UFPA: Belém e "
        "Castanhal, ambas Conceito Enade 4. Não existe oferta UFPA Conceito 1 e nenhum grupo "
        "artificial é criado. Os arquivos temáticos são agregados por CO_CURSO antes das "
        "junções. Belém apresenta maior participação e NT_GER médio; Castanhal apresenta "
        "percepções mais favoráveis do processo formativo e recomendação do curso. As duas "
        "ofertas ficam abaixo das medianas de seus benchmarks estruturais comparáveis. "
        "As interpretações são descritivas e não causais."
    )
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.add_run(
        "Palavras-chave: Enade; Educação Física; UFPA; formação de professores; microdados; benchmark."
    )

    doc.add_heading("1 INTRODUÇÃO", level=1)
    _adicionar_texto(
        doc,
        "A pergunta central é: quais características de desempenho, participação, composição "
        "discente, trajetória acadêmica e avaliação do processo formativo caracterizam as "
        "ofertas de Educação Física da UFPA e como elas se posicionam em relação às demais "
        "ofertas da mesma área no Pará, na Região Norte e no Brasil?",
    )
    _adicionar_texto(
        doc,
        "Belém e Castanhal possuem Conceito Enade 4. O relatório privilegia a heterogeneidade "
        "interna e a posição relativa perante referências territoriais e estruturais, sem "
        "converter Conceito 4 em categoria de insuficiência.",
    )

    doc.add_heading("2 REFERENCIAL INSTITUCIONAL E METODOLÓGICO", level=1)
    _adicionar_texto(
        doc,
        "O Conceito Enade é tratado como classificação externa do curso e não como variável "
        "causal. Diferenças entre ofertas e benchmarks são apresentadas como padrões e "
        "hipóteses para investigação institucional.",
    )

    doc.add_heading("3 METODOLOGIA", level=1)
    for texto in (
        "A unidade principal é CO_CURSO. Não se usa posição de linha como chave, não se cria "
        "identificador artificial e não se realizam joins individuais entre arquivos temáticos. "
        "O fluxo é arquivo temático → tratamento de ausências → agregação por CO_CURSO → uma "
        "linha por curso → junções one-to-one → comparação entre cursos.",
        "NT_GER, NT_OBJ e NT_DIS podem ser examinadas conjuntamente porque pertencem ao mesmo "
        "arquivo. Relações entre desempenho e indicadores de outros arquivos são apenas ecológicas.",
        "O Grupo A (UFPA Conceito 1) possui N=0. Os grupos exclusivos efetivos são UFPA com "
        "conceito superior, outras IES do Pará, restante da Região Norte e restante do Brasil.",
        "O benchmark principal mantém modalidade, categoria administrativa e organização "
        "acadêmica, com participantes entre 0,5x e 2x o porte da oferta UFPA.",
        "QE_I20–QE_I66 utilizam respostas 1–6; 7 e 8 são ausências analíticas. QE_I31, QE_I32 "
        "e QE_I43 são específicos de EaD. Não foi criado índice global do processo formativo.",
    ):
        _adicionar_texto(doc, texto)

    doc.add_heading("4 PANORAMA DA LICENCIATURA EM EDUCAÇÃO FÍSICA", level=1)
    _adicionar_texto(
        doc,
        "O universo analítico reúne 406 cursos. A UFPA possui duas ofertas localizadas nas "
        "fontes oficiais: Belém (CO_CURSO 104598) e Castanhal (CO_CURSO 21849), ambas "
        "presenciais e com Conceito Enade 4.",
    )
    adicionar_tabela(doc, "Tabela 1 – Ofertas de Educação Física da UFPA", tabela_ofertas(base), FONTE_DADOS)
    adicionar_figura(
        doc,
        figuras / "01_painel_ofertas_ufpa.png",
        "Figura 1 – Participação nas ofertas da UFPA",
        FONTE_DADOS,
    )

    doc.add_heading("5 RESULTADOS", level=1)
    doc.add_heading("5.1 Desempenho", level=2)
    _adicionar_texto(
        doc,
        f"Belém apresenta NT_GER médio {_fmt(bel['nt_ger_mean'])}, mediana "
        f"{_fmt(bel['nt_ger_median'])} e participação {_pct(bel['TAXA_PARTICIPACAO_OFICIAL'])}. "
        f"Castanhal apresenta NT_GER médio {_fmt(cas['nt_ger_mean'])}, mediana "
        f"{_fmt(cas['nt_ger_median'])} e participação {_pct(cas['TAXA_PARTICIPACAO_OFICIAL'])}. "
        f"A diferença Belém–Castanhal em NT_GER é {_fmt(bel['nt_ger_mean'] - cas['nt_ger_mean'])} pontos.",
    )
    _adicionar_texto(
        doc,
        f"Em NT_OBJ, as médias são {_fmt(bel['nt_obj_mean'])} em Belém e "
        f"{_fmt(cas['nt_obj_mean'])} em Castanhal. Em NT_DIS, Castanhal registra "
        f"{_fmt(cas['nt_dis_mean'])} e Belém {_fmt(bel['nt_dis_mean'])}.",
    )
    adicionar_figura(doc, figuras / "02_posicao_relativa_nt_ger.png", "Figura 2 – Posição relativa em NT_GER", FONTE_DADOS)
    adicionar_figura(doc, figuras / "03a_distribuicao_nt_ger.png", "Figura 3 – Distribuição individual de NT_GER", FONTE_DADOS)
    adicionar_figura(doc, figuras / "03b_distribuicao_nt_obj.png", "Figura 4 – Distribuição individual de NT_OBJ", FONTE_DADOS)
    adicionar_figura(doc, figuras / "03c_distribuicao_nt_dis.png", "Figura 5 – Distribuição individual de NT_DIS", FONTE_DADOS)

    doc.add_heading("5.2 Perfil demográfico e socioeconômico", level=2)
    _adicionar_texto(
        doc,
        f"A renda de até três salários mínimos alcança {_pct(bel['renda_ate_3sm_pct'])} em Belém "
        f"e {_pct(cas['renda_ate_3sm_pct'])} em Castanhal. Trabalham {_pct(bel['trabalha_pct'])} "
        f"e {_pct(cas['trabalha_pct'])}, respectivamente. A maior diferença selecionada ocorre "
        f"em bolsa acadêmica: {_pct(bel['bolsa_academica_pct'])} em Belém e "
        f"{_pct(cas['bolsa_academica_pct'])} em Castanhal.",
    )
    adicionar_tabela(doc, "Tabela 2 – Perfil selecionado das ofertas UFPA", tabela_perfil(base), FONTE_DADOS)
    adicionar_figura(doc, figuras / "04_perfil_socioeconomico.png", "Figura 6 – Perfil socioeconômico", FONTE_DADOS)

    doc.add_heading("5.3 Trajetória e condições acadêmicas", level=2)
    _adicionar_texto(
        doc,
        f"O tempo médio desde o ingresso é {_fmt(bel['anos_desde_ingresso_media'])} anos em Belém "
        f"e {_fmt(cas['anos_desde_ingresso_media'])} em Castanhal. A proporção que estuda quatro "
        f"horas ou mais por semana é {_pct(bel['estudo_4h_ou_mais_pct'])} e "
        f"{_pct(cas['estudo_4h_ou_mais_pct'])}. A intenção de atuar no magistério é "
        f"{_pct(bel['pretende_magisterio_pct'])} e {_pct(cas['pretende_magisterio_pct'])}.",
    )

    doc.add_heading("5.4 Processo formativo", level=2)
    _adicionar_texto(
        doc,
        "As oito dimensões exploratórias apresentam consistência interna elevada no recorte "
        "nacional e na UFPA. Alfa elevado não comprova unidimensionalidade e os escores não "
        "formam um índice global.",
    )
    adicionar_tabela(doc, "Tabela 3 – Dimensões do processo formativo", tabela_dimensoes(base), FONTE_DADOS)
    _adicionar_texto(
        doc,
        "Castanhal apresenta médias superiores nas oito dimensões, enquanto Belém apresenta "
        "NT_GER e NT_OBJ médios maiores. Desempenho e percepção formativa são dimensões distintas.",
    )
    adicionar_figura(doc, figuras / "05_processo_formativo_dimensoes.png", "Figura 7 – Processo formativo", FONTE_DADOS)

    doc.add_heading("5.5 Recomendação", level=2)
    _adicionar_texto(
        doc,
        f"No QE_I68, recomendação do curso, Castanhal apresenta média {_fmt(cas['qe_i68_media'])} "
        f"e Belém {_fmt(bel['qe_i68_media'])}. No QE_I69, recomendação da IES, as médias são "
        f"{_fmt(cas['qe_i69_media'])} e {_fmt(bel['qe_i69_media'])}. Os itens não são denominados "
        "automaticamente como satisfação.",
    )
    adicionar_figura(doc, figuras / "07_recomendacao.png", "Figura 8 – Recomendação do curso e da IES", FONTE_DADOS)

    doc.add_heading("5.6 Benchmark comparável", level=2)
    tab_b = tabela_benchmark(benchmark)
    adicionar_tabela(doc, "Tabela 4 – Benchmark comparável de NT_GER", tab_b, FONTE_DADOS)
    linha_c = tab_b.loc[tab_b["Oferta"].str.contains("Castanhal")].iloc[0]
    linha_b = tab_b.loc[tab_b["Oferta"].str.contains("Belém")].iloc[0]
    _adicionar_texto(
        doc,
        f"Castanhal fica {_fmt(linha_c['Dif. mediana'])} pontos abaixo da mediana de 9 cursos "
        f"comparáveis; Belém, {_fmt(linha_b['Dif. mediana'])} pontos abaixo da mediana de 24. "
        "A janela mais estrita de Castanhal contém apenas um comparável e é considerada frágil.",
    )
    adicionar_figura(doc, figuras / "06_benchmark_sensibilidade.png", "Figura 9 – Sensibilidade do benchmark", FONTE_DADOS)

    doc.add_heading("5.7 Associações ecológicas", level=2)
    tab_a = associacoes[["X", "N_CURSOS", "SPEARMAN_RHO", "P_VALOR"]].copy()
    adicionar_tabela(doc, "Tabela 5 – Associações ecológicas com NT_GER", tab_a, FONTE_DADOS)
    _adicionar_texto(
        doc,
        "A maior magnitude entre os pares examinados é a associação positiva entre auxílio "
        "permanência e NT_GER médio (rho=0,437; N=347). Trabalho apresenta associação negativa "
        "(rho=-0,290). Essas relações são agregadas por curso e não permitem inferência individual ou causal.",
    )

    doc.add_heading("5.8 Síntese gráfica", level=2)
    adicionar_figura(doc, figuras / "08_sintese_ufpa.png", "Figura 10 – Síntese descritiva das ofertas UFPA", FONTE_DADOS)

    doc.add_heading("6 DISCUSSÃO", level=1)
    for texto in (
        "Belém combina maior taxa de participação, maior NT_GER e maior NT_OBJ, além de posição "
        "relativa mais alta. Castanhal apresenta percepções mais favoráveis do processo formativo "
        "e recomendação do curso muito superior. O contraste mostra que a avaliação é multidimensional.",
        "As comparações territoriais amplas posicionam as duas ofertas acima das médias externas, "
        "mas ambas ficam abaixo das medianas de benchmarks estruturais comparáveis. Posição favorável "
        "no universo total não implica vantagem frente a pares semelhantes.",
        "Diferenças em bolsa acadêmica, escolaridade parental, trabalho e auxílio permanência são "
        "hipóteses institucionais para aprofundamento e não mecanismos causais demonstrados.",
    ):
        _adicionar_texto(doc, texto)

    doc.add_heading("7 CONCLUSÃO", level=1)
    _adicionar_texto(
        doc,
        "As duas ofertas da UFPA possuem Conceito Enade 4 e se posicionam acima da maior parte "
        "dos cursos nacionais em NT_GER. Belém apresenta maior participação e desempenho médio, "
        "sobretudo objetivo; Castanhal apresenta percepções formativas e recomendação do curso "
        "mais favoráveis. Ambas ficam abaixo das medianas dos respectivos benchmarks comparáveis. "
        "O conjunto não permite atribuir causalidade às diferenças observadas.",
    )

    doc.add_heading("REFERÊNCIAS", level=1)
    adicionar_referencias(doc)

    doc.add_heading("APÊNDICES", level=1)
    doc.add_heading("APÊNDICE A – REGRAS DE INTEGRIDADE", level=2)
    for texto in (
        "Unidade principal CO_CURSO; nenhum join individual entre arquivos temáticos.",
        "Ausência de conceito não é Conceito 1 e o Grupo A não é reconstruído.",
        "Associações entre temas distintos são exclusivamente ecológicas.",
        "QE_I68 e QE_I69 são tratados como recomendação, não satisfação.",
    ):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.first_line_indent = Cm(0)
        p.add_run(texto)

    doc.add_heading("APÊNDICE B – APROFUNDAMENTOS SUGERIDOS", level=2)
    aprofundamentos = (
        "Desempenho por componente e distribuição — aprofundar NT_OBJ, NT_DIS, QT_ACERTOS e PROFICIÊNCIA no mesmo arquivo, com ECDF, quantis e tamanhos de efeito; limitação: relações mecânicas entre indicadores.",
        "Oportunidades acadêmicas e bolsas — investigar a diferença de bolsa acadêmica entre campi com referências federais comparáveis; limitação: ausência de identificação individual entre temas.",
        "Processo formativo item a item — decompor as dimensões com maiores diferenças usando QE_I20–QE_I66 e N válido; limitação: respostas autorreferidas e múltiplas comparações.",
        "Benchmark estrutural mais estrito — aplicar matching ecológico multivariado e sensibilidade; limitação: confundimento residual e pequeno N em estratos restritos.",
        "Recomendação e processo formativo — avaliar associações ecológicas entre QE_I68/QE_I69 e dimensões formativas no universo nacional; limitação: falácia ecológica.",
    )
    for texto in aprofundamentos:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.first_line_indent = Cm(0)
        p.add_run(texto)

    destino.parent.mkdir(parents=True, exist_ok=True)
    doc.save(destino)


def gerar_markdown(produtos: dict[str, pd.DataFrame], destino: Path) -> None:
    base = produtos["base"]
    bel = _oferta(base, 104598)
    cas = _oferta(base, 21849)
    texto = f"""# EDUCAÇÃO FÍSICA NO ENADE DAS LICENCIATURAS 2025: DESEMPENHO, PERFIL, PROCESSO FORMATIVO E BENCHMARKS DAS OFERTAS DA UFPA

## RESUMO

O universo analítico reúne 406 cursos de Educação Física, com duas ofertas UFPA: Belém e Castanhal, ambas presenciais e Conceito Enade 4. Não existe oferta UFPA Conceito 1. Os arquivos temáticos foram agregados por CO_CURSO antes de qualquer junção. Belém apresenta participação de {_pct(bel['TAXA_PARTICIPACAO_OFICIAL'])} e NT_GER médio {_fmt(bel['nt_ger_mean'])}; Castanhal, {_pct(cas['TAXA_PARTICIPACAO_OFICIAL'])} e {_fmt(cas['nt_ger_mean'])}. Castanhal apresenta percepções mais favoráveis do processo formativo e maior recomendação do curso. As duas ofertas ficam abaixo das medianas de benchmarks estruturais comparáveis. As interpretações são descritivas e não causais.

**Palavras-chave:** Enade; Educação Física; UFPA; formação de professores; microdados; benchmark.

# 1 INTRODUÇÃO

A pergunta central é: quais características de desempenho, participação, composição discente, trajetória acadêmica e avaliação do processo formativo caracterizam as ofertas de Educação Física da UFPA e como elas se posicionam em relação às demais ofertas da mesma área no Pará, na Região Norte e no Brasil?

# 2 REFERENCIAL INSTITUCIONAL E METODOLÓGICO

O Conceito Enade é tratado como classificação externa do curso. Não existe oferta da UFPA com Conceito Enade 1 em Educação Física e nenhum grupo artificial é criado.

# 3 METODOLOGIA

A unidade principal é CO_CURSO. Não são realizados joins individuais entre arquivos temáticos; as integrações ocorrem somente depois da agregação por curso e validação one-to-one. Relações entre temas distintos são ecológicas.

# 4 PANORAMA DA LICENCIATURA EM EDUCAÇÃO FÍSICA

Foram identificados 406 cursos, com duas ofertas UFPA: Belém (CO_CURSO 104598) e Castanhal (CO_CURSO 21849), ambas Conceito Enade 4.

# 5 RESULTADOS

## 5.1 Desempenho

Belém apresenta NT_GER médio {_fmt(bel['nt_ger_mean'])}; Castanhal, {_fmt(cas['nt_ger_mean'])}. Em NT_OBJ, as médias são {_fmt(bel['nt_obj_mean'])} e {_fmt(cas['nt_obj_mean'])}. Em NT_DIS, {_fmt(bel['nt_dis_mean'])} e {_fmt(cas['nt_dis_mean'])}.

## 5.2 Perfil demográfico e socioeconômico

Renda até 3 SM: Belém {_pct(bel['renda_ate_3sm_pct'])}; Castanhal {_pct(cas['renda_ate_3sm_pct'])}. Bolsa acadêmica: {_pct(bel['bolsa_academica_pct'])} e {_pct(cas['bolsa_academica_pct'])}.

## 5.3 Trajetória e condições acadêmicas

Tempo médio desde ingresso: Belém {_fmt(bel['anos_desde_ingresso_media'])}; Castanhal {_fmt(cas['anos_desde_ingresso_media'])}. Pretende magistério: {_pct(bel['pretende_magisterio_pct'])} e {_pct(cas['pretende_magisterio_pct'])}.

## 5.4 Processo formativo

Castanhal apresenta médias superiores nas oito dimensões exploratórias. Os escores não constituem índice global e alfa elevado não demonstra unidimensionalidade.

## 5.5 Recomendação

QE_I68: Castanhal {_fmt(cas['qe_i68_media'])}; Belém {_fmt(bel['qe_i68_media'])}. QE_I69: {_fmt(cas['qe_i69_media'])} e {_fmt(bel['qe_i69_media'])}.

## 5.6 Benchmark comparável

Castanhal possui 9 comparáveis no cenário principal e Belém 24. Ambas ficam abaixo da mediana comparável de NT_GER.

## 5.7 Associações ecológicas

As associações são calculadas entre cursos e não permitem inferência individual ou causal. A maior magnitude examinada é auxílio permanência × NT_GER médio (rho=0,437; N=347).

# 6 DISCUSSÃO

Belém se destaca em participação e desempenho; Castanhal em processo formativo percebido e recomendação. O contraste evidencia multidimensionalidade e não sustenta uma explicação causal única.

# 7 CONCLUSÃO

As duas ofertas possuem Conceito Enade 4, apresentam posição nacional relativamente favorável em NT_GER e perfis internos distintos. A interpretação institucional deve combinar desempenho, perfil, trajetória, processo formativo, recomendação e benchmarks comparáveis.

# REFERÊNCIAS

As referências normativas e bibliográficas são inseridas integralmente na versão DOCX pelo módulo compartilhado de referências do projeto.

# APÊNDICES

Aprofundamentos sugeridos: componentes do desempenho; oportunidades acadêmicas e bolsas; processo formativo item a item; matching ecológico mais estrito; recomendação e processo formativo.
"""
    destino.parent.mkdir(parents=True, exist_ok=True)
    destino.write_text(texto, encoding="utf-8")


def gerar_relatorio(root: Path) -> ResultadoRelatorio:
    produtos = carregar_produtos(root)
    pasta = root / "relatorios" / "educacao_fisica"
    docx = pasta / "relatorio_educacao_fisica_enade_2025_ufpa.docx"
    md = pasta / "relatorio_educacao_fisica_enade_2025_ufpa.md"
    gerar_docx(root, produtos, docx)
    gerar_markdown(produtos, md)
    conversao = converter_docx_para_pdf(docx, pasta)
    return ResultadoRelatorio(
        docx=docx,
        markdown=md,
        conversao_pdf=conversao,
    )
